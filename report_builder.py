"""
주간 수급 분석 엑셀 생성기.

노션 '주간 수급 데이터' DB(또는 같은 열 구성의 엑셀/CSV)를 받아
  - '주간 수급 데이터' 시트 : 원본 그대로
  - '수급분석' 시트          : 원본을 참조하는 수식 기반 피벗 4종 + 차트 2개 + 자동 요약
을 만든다. 2026-10-09 수작업으로 만든 '26 10월 둘째주 수급누적.xlsx'와 같은 구성이다.
"""

from __future__ import annotations

from datetime import date

import openpyxl
import pandas as pd
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.data_source import StrRef
from openpyxl.chart.series import SeriesLabel
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAW_SHEET = "주간 수급 데이터"
OUT_SHEET = "수급분석"
THEME_SHEET = "테마분석"
INV = ["외국인", "기관합계", "연기금"]
# 노션 내보내기와 같은 열 순서 (수식이 열 위치가 아니라 이름으로 찾으므로 순서가 바뀌어도 된다)
RAW_COLUMNS = ["종목명", "기준주간", "뉴스", "레코드키", "리포트", "비고", "섹터", "수집시각",
               "순매수금액", "순위", "시장", "연속순매수일", "원칙 점검", "종목코드", "종목페이지",
               "주간등락률", "주차", "차트", "투자주체"]
WINDOW = 6          # 누적 분석에 쓰는 최근 주 수
TOP_N = 20          # 누적 순위 표 종목 수
PULLBACK_STREAK = 5  # 눌림매집: 연속순매수일 이 값 이상 + 주간 하락
LONG_STREAK = 8      # 장기연속 표시 기준

FONT = "맑은 고딕"
F = {
    "base": Font(name=FONT, size=10),
    "b": Font(name=FONT, size=10, bold=True),
    "h": Font(name=FONT, size=10, bold=True, color="FFFFFF"),
    "t": Font(name=FONT, size=15, bold=True),
    "sec": Font(name=FONT, size=12, bold=True, color="1F3864"),
    "note": Font(name=FONT, size=9, color="595959", italic=True),
}
FILL_H = PatternFill("solid", fgColor="1F3864")
FILL_TOT = PatternFill("solid", fgColor="E7ECF5")
FILL_BOX = PatternFill("solid", fgColor="F7F7F2")
_thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)
NUM = '#,##0;[Red]-#,##0;"-"'
PCT = '0.00;[Red]-0.00;"0.00"'
CENTER = Alignment(horizontal="center", vertical="center")


# --- 이름 규칙 -----------------------------------------------------------------

def week_label_to_monday(label: str) -> date:
    year, week = label.split("-W")
    return date.fromisocalendar(int(year), int(week), 1)


def report_filename(label: str) -> str:
    """'2026-W41' -> '26 10월 둘째주 수급누적.xlsx'.

    그 주의 목요일이 속한 달의 몇째 주인지로 센다(1일이 목요일 이전에 끼면 그 주가 첫째 주).
    """
    thursday = date.fromisocalendar(*map(int, label.split("-W")), 4)
    nth = (thursday.day - 1) // 7 + 1
    names = ["첫째", "둘째", "셋째", "넷째", "다섯째"]
    return f"{thursday.year % 100:02d} {thursday.month}월 {names[nth - 1]}주 수급누적.xlsx"


def week_title(label: str) -> str:
    return report_filename(label).replace(" 수급누적.xlsx", "")


# --- 데이터 정리 -----------------------------------------------------------------

def normalize(df: pd.DataFrame) -> pd.DataFrame:
    df = df.copy()
    for c in RAW_COLUMNS:
        if c not in df.columns:
            df[c] = None
    df = df[RAW_COLUMNS]
    df = df[df["주차"].notna() & df["투자주체"].isin(INV)]
    for c in ["순매수금액", "순위", "연속순매수일", "주간등락률"]:
        df[c] = pd.to_numeric(df[c], errors="coerce")
    df["연속순매수일"] = df["연속순매수일"].fillna(0)
    df["주간등락률"] = df["주간등락률"].fillna(0)
    df = df.sort_values(["주차", "투자주체", "순위"], ascending=[False, True, True])
    return df.reset_index(drop=True)


# --- 자동 요약 ------------------------------------------------------------------

def _fmt_eok(v: float) -> str:
    return f"{v / 10000:.1f}조" if abs(v) >= 10000 else f"{v:,.0f}억"


def build_insights(df: pd.DataFrame) -> list[str]:
    weeks = sorted(df["주차"].unique())
    cur = weeks[-1]
    prev = weeks[-2] if len(weeks) > 1 else None
    win = weeks[-WINDOW:]
    wdf = df[df["주차"].isin(win)]
    cdf = df[df["주차"] == cur]
    lines = []

    tot = wdf.pivot_table(index="종목명", columns="투자주체", values="순매수금액", aggfunc="sum", fill_value=0)
    tot["합계"] = tot.sum(axis=1)
    tot = tot.sort_values("합계", ascending=False)
    top = tot.index[0]
    parts = " + ".join(f"{inv} {_fmt_eok(tot.loc[top, inv])}" for inv in INV if inv in tot and tot.loc[top, inv] > 0)
    runner = ", ".join(tot.index[1:3])
    lines.append(f"① 최근 {len(win)}주 누적 1위는 {top}({_fmt_eok(tot.loc[top, '합계'])}: {parts}). 2·3위는 {runner}.")

    g = cdf.groupby("종목명")["투자주체"].nunique()
    three = g[g >= 3].index.tolist()
    two = g[g == 2].index.tolist()
    lines.append("② 이번 주 3주체 동시 매수: " + (", ".join(three) if three else "없음")
                 + (" → 외국인·기관·연기금이 모두 같은 종목을 샀습니다." if three else " → 세 주체의 시선이 갈렸습니다."))

    streak = cdf.sort_values("연속순매수일", ascending=False).head(3)
    s_txt = ", ".join(f"{r.종목명}({r.투자주체} {int(r.연속순매수일)}일)" for r in streak.itertuples())
    lines.append(f"③ 쌍끌이(2주체): {', '.join(two) if two else '없음'} · 가장 끈질긴 연속 순매수: {s_txt}.")

    agg = cdf.groupby("종목명").agg(연속=("연속순매수일", "max"), 등락=("주간등락률", "mean"))
    pull = agg[(agg["연속"] >= PULLBACK_STREAK) & (agg["등락"] < 0)]
    if len(pull):
        p_txt = ", ".join(f"{n}(연속 {int(r.연속)}일, 주간 {r.등락:+.1f}%)" for n, r in pull.iterrows())
        lines.append(f"④ 눌림매집 후보: {p_txt}. 사는 주체는 버티는데 주가가 쉬는 구간입니다.")
    else:
        lines.append(f"④ 눌림매집 후보 없음 (연속 {PULLBACK_STREAK}일 이상 순매수 + 주간 하락 종목이 없습니다).")

    wk = df[df["주차"].isin(weeks[-3:])].pivot_table(index="주차", columns="투자주체", values="순매수금액", aggfunc="sum")
    trend = []
    for inv in ["외국인", "기관합계"]:
        if inv in wk:
            seq = [wk.loc[w, inv] for w in wk.index]
            arrow = " → ".join(_fmt_eok(v) for v in seq)
            mood = "강해지는 중" if seq[-1] > seq[0] else "식는 중"
            trend.append(f"{inv} {arrow}({mood})")
    lines.append("⑤ 매수 강도 추이(상위10 합, 최근 3주): " + " · ".join(trend) + ".")

    if prev:
        new = cdf[~cdf["종목명"].isin(df[df["주차"] == prev]["종목명"])]
        new = new.groupby("종목명")["순매수금액"].sum().sort_values(ascending=False).head(3)
        if len(new):
            lines.append("⑥ 이번 주 새로 등장한 큰 매수: " + ", ".join(f"{n}({_fmt_eok(v)})" for n, v in new.items()) + ".")
    return lines


# --- 엑셀 작성 ------------------------------------------------------------------

class _Sheet:
    def __init__(self, ws):
        self.ws = ws

    def put(self, cell, value, font="base", fmt=None, fill=None, align=None, border=True):
        c = self.ws[cell]
        c.value = value
        c.font = F[font]
        if fmt:
            c.number_format = fmt
        if fill:
            c.fill = fill
        if align:
            c.alignment = align
        if border:
            c.border = BOX
        return c

    def header(self, row, start_col, labels):
        for i, label in enumerate(labels):
            self.put(f"{get_column_letter(start_col + i)}{row}", label, "h", fill=FILL_H,
                     align=Alignment(horizontal="center", vertical="center", wrap_text=True))
        self.ws.row_dimensions[row].height = 30

    def section(self, row, text, note):
        self.put(f"A{row}", text, "sec", border=False)
        self.put(f"A{row + 1}", note, "note", border=False)


def write_raw(wb, df: pd.DataFrame):
    ws = wb.create_sheet(RAW_SHEET)
    ws.append(RAW_COLUMNS)
    for row in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and pd.isna(v)) else v for v in row])
    for c in ws[1]:
        c.font = F["b"]
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = F["base"]
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    for col, w in {"A": 16, "I": 11, "Q": 10, "S": 9}.items():
        ws.column_dimensions[col].width = w
    return ws


def write_theme_sheet(wb, themes: dict, cur: str, n_win: int) -> None:
    """'테마분석' 시트: 이번 주 테마 순위 · 최근 n주 누적 테마 · 종목별 대표 테마 (값으로 기록)."""
    from theme_analysis import commentary

    ws = wb.create_sheet(THEME_SHEET, 1)
    s = _Sheet(ws)
    s.put("A1", "테마별 수급 분석", "t", border=False)
    s.put("A2", f"출처: {themes['source']} · 수급 종목이 {2}개 이상 모인 테마만 · 한 종목이 여러 테마에 속하면 모두 포함"
                "(테마 합계끼리 중복 가능) · 구성 종목이 같은 테마는 ' / '로 묶음 · 금액 억원", "note", border=False)
    s.put("A3", "상태: 새로 부상 = 직전 3주 없음 → 이번 주 등장 · 가속/둔화 = 이번 주가 직전 3주 평균의 1.5배 이상/0.5배 이하",
          "note", border=False)
    r = 5
    for line in themes["insights"]:
        s.put(f"A{r}", "· " + line, fill=FILL_BOX, border=False)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=10)
        r += 1

    def table(start, title, note, items, with_cur):
        s.section(start, title, note)
        heads = ["테마", "종목수", "외국인", "기관합계", "연기금", "합계"] + (["이번 주"] if with_cur else []) \
            + ["상태", "등장 주수", "구성 종목 (금액순)"]
        s.header(start + 2, 1, heads)
        row = start + 3
        for t in items:
            vals = [t["테마"], t["종목수"], t["외국인"], t["기관합계"], t["연기금"], t["합계"]] \
                + ([t["이번주"]] if with_cur else []) + [t["상태"], t["등장주수"], ", ".join(t["종목"])]
            for j, v in enumerate(vals):
                cl = get_column_letter(1 + j)
                is_num = isinstance(v, float)
                s.put(f"{cl}{row}", v, "b" if j == 0 else "base", NUM if is_num else None,
                      align=CENTER if j in (1,) or heads[j] in ("상태", "등장 주수") else None)
            row += 1
        return start + 3, row - 1, len(heads)

    if themes["current"]:
        f1, l1, _ = table(r + 1, f"1. 이번 주({cur}) 수급이 모인 테마", "합계 큰 순",
                          themes["current"][:20], with_cur=False)
        ch = BarChart()
        ch.type, ch.grouping, ch.overlap = "bar", "stacked", 100
        ch.title = "이번 주 테마별 순매수 (억원)"
        n = min(10, l1 - f1 + 1)
        for j in range(3):
            ser = Series(Reference(ws, min_col=3 + j, min_row=f1, max_row=f1 + n - 1))
            ser.tx = SeriesLabel(strRef=StrRef(f"'{THEME_SHEET}'!${get_column_letter(3 + j)}${f1 - 1}"))
            ch.series.append(ser)
        ch.set_categories(Reference(ws, min_col=1, min_row=f1, max_row=f1 + n - 1))
        ch.x_axis.scaling.orientation = "maxMin"
        ch.x_axis.tickLblSkip = 1
        ch.y_axis.numFmt = "#,##0"
        ch.legend.position = "b"
        ch.height, ch.width = 9, 16
        ch.x_axis.delete = ch.y_axis.delete = False
        ws.add_chart(ch, f"L{r + 1}")
        r = max(l1, f1 + 18) + 3
    else:
        s.put(f"A{r + 1}", "이번 주에는 수급 종목이 2개 이상 모인 테마가 없습니다.", "note", border=False)
        r += 3
    if themes["cumulative"]:
        _, l2, _ = table(r, f"2. 최근 {n_win}주 누적 테마 (이번 주와 비교)", "누적 합계 큰 순. '이번 주'가 0이면 이번 주에 빠진 테마",
                         themes["cumulative"][:25], with_cur=True)
        r = l2 + 3
        s.section(r, "3. 테마 한 줄 해설 (누적 상위 10)", "규칙 기반 자동 문장")
        for i, t in enumerate(themes["cumulative"][:10]):
            s.put(f"A{r + 2 + i}", f"{t['테마']}", "b")
            s.put(f"B{r + 2 + i}", commentary(t))
            ws.merge_cells(start_row=r + 2 + i, start_column=2, end_row=r + 2 + i, end_column=11)
        r = r + 2 + min(10, len(themes["cumulative"])) + 2
    if themes["stock_themes"]:
        s.section(r, "4. 이번 주 수급 종목의 대표 테마", "수급 종목이 모인 테마 중 최대 2개")
        s.header(r + 2, 1, ["종목명", "대표 테마"])
        for i, (name, ts) in enumerate(themes["stock_themes"].items()):
            s.put(f"A{r + 3 + i}", name, "b")
            s.put(f"B{r + 3 + i}", ", ".join(ts) or "-")
            ws.merge_cells(start_row=r + 3 + i, start_column=2, end_row=r + 3 + i, end_column=6)
    widths = {"A": 30, "B": 9, "C": 11, "D": 11, "E": 11, "F": 11, "G": 11, "H": 11, "I": 10, "J": 50}
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A4"


def build_workbook(df: pd.DataFrame, path: str, themes: dict | None = None) -> dict:
    """엑셀을 만들어 path에 저장하고, 노션에 쓸 요약 정보를 돌려준다.

    themes: theme_analysis.analyze() 결과. 주면 '테마분석' 시트와 요약 한 줄이 추가된다.
    """
    df = normalize(df)
    weeks = sorted(df["주차"].unique())
    if not weeks:
        raise ValueError("투자주체(외국인·기관합계·연기금) 데이터가 없습니다.")
    cur = weeks[-1]
    prev = weeks[-2] if len(weeks) > 1 else weeks[-1]
    win = weeks[-WINDOW:]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(OUT_SHEET)
    write_raw(wb, df)
    s = _Sheet(ws)

    end = len(df) + 1
    S = f"'{RAW_SHEET}'!"
    col = {name: get_column_letter(RAW_COLUMNS.index(name) + 1) for name in RAW_COLUMNS}
    R = {k: f"{S}${col[n]}$2:${col[n]}${end}" for k, n in
         {"종목": "종목명", "금액": "순매수금액", "시장": "시장", "연속": "연속순매수일",
          "등락": "주간등락률", "주차": "주차", "주체": "투자주체"}.items()}

    # 제목·요약
    s.put("A1", "주간 수급 분석 — 외국인 · 기관합계 · 연기금", "t", border=False)
    s.put("A2", f"원본: '{RAW_SHEET}' 시트 ({weeks[0]} ~ {cur}, 주체별 주간 순매수 상위 10종목). 금액 단위 억원, 등락률 단위 %. "
                "모든 표는 원본 시트를 참조하는 수식입니다.", "note", border=False)
    s.put("A3", "※ 원본이 '주체별 상위 10종목'만 담고 있어 합계는 시장 전체 순매수가 아니라 '상위 10종목 매수 강도'로 읽어야 합니다. "
                "휴장일이 낀 주는 거래일이 적습니다.", "note", border=False)
    insights = build_insights(df)
    if themes and themes.get("current"):
        top = themes["current"][:3]
        insights.append(f"⑦ 테마: 이번 주 수급이 가장 몰린 테마는 "
                        + ", ".join(f"{t['테마'].split(' / ')[0]}({t['종목수']}종목)" for t in top)
                        + f" → 자세한 내용은 '{THEME_SHEET}' 시트.")
    s.put("A5", f"핵심 요약 ({date.today().isoformat()} 자동 작성, {week_title(cur)} 기준)", "b", fill=FILL_BOX, border=False)
    ws.merge_cells("A5:L5")
    for i, text in enumerate(insights):
        r = 6 + i
        s.put(f"A{r}", text, fill=FILL_BOX, border=False)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=12)

    # 표 1: 주차 × 투자주체
    r0 = 6 + len(insights) + 2
    s.section(r0, "1. 주차별 · 투자주체별 순매수 합계 (억원)", "상위 10종목 순매수 합. 주체별 매수 강도의 주간 추이")
    s.header(r0 + 2, 1, ["주차"] + INV + ["합계"])
    w_first = r0 + 3
    for i, w in enumerate(weeks):
        r = w_first + i
        s.put(f"A{r}", w, align=CENTER)
        for j in range(3):
            cl = get_column_letter(2 + j)
            s.put(f"{cl}{r}", f'=SUMIFS({R["금액"]},{R["주차"]},$A{r},{R["주체"]},{cl}${r0 + 2})', fmt=NUM)
        s.put(f"E{r}", f"=SUM(B{r}:D{r})", "b", NUM)
    w_last = w_first + len(weeks) - 1
    rt = w_last + 1
    s.put(f"A{rt}", "전체 합계", "b", fill=FILL_TOT, align=CENTER)
    for cl in "BCDE":
        s.put(f"{cl}{rt}", f"=SUM({cl}{w_first}:{cl}{w_last})", "b", NUM, FILL_TOT)
    WIN = f"$A${w_last - len(win) + 1}:$A${w_last}"  # 최근 WINDOW주 범위

    ch = BarChart()
    ch.type, ch.grouping = "col", "clustered"
    ch.title = "주차별 순매수 (억원)"
    ch.y_axis.numFmt = "#,##0"
    ch.legend.position = "b"
    ch.add_data(Reference(ws, min_col=2, max_col=4, min_row=r0 + 2, max_row=w_last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=w_first, max_row=w_last))
    ch.height, ch.width = 7.5, 16
    ch.x_axis.delete = ch.y_axis.delete = False
    ws.add_chart(ch, f"G{r0}")

    # 표 2: 이번 주 신호판
    cdf = df[df["주차"] == cur]
    order = cdf.groupby("종목명")["순매수금액"].sum().sort_values(ascending=False).index.tolist()
    r1 = max(rt, r0 + 16) + 3
    s.section(r1, f"2. 이번 주({cur}) 수급 신호판",
              f"참여주체 3 = 3주체 동시, 2 = 쌍끌이 · 직전주 미등장 = 신규 · 연속 {PULLBACK_STREAK}일 이상인데 주가 하락 = 눌림매집")
    s.header(r1 + 2, 1, ["종목명", "시장", "외국인", "기관합계", "연기금", "합계", "참여주체수",
                         "최대 연속순매수일", "주간등락률(%)", "직전주 등장", "신호"])
    s.put(f"N{r1 + 2}", cur, "note", border=False)
    s.put(f"N{r1 + 3}", prev, "note", border=False)
    s.put(f"O{r1 + 2}", "← 이번 주", "note", border=False)
    s.put(f"O{r1 + 3}", "← 직전 주 (값을 바꾸면 다른 주로 다시 계산)", "note", border=False)
    CW, PW = f"$N${r1 + 2}", f"$N${r1 + 3}"
    s_first = r1 + 3
    for i, name in enumerate(order):
        r = s_first + i
        s.put(f"A{r}", name, "b")
        s.put(f"B{r}", f'=INDEX({R["시장"]},MATCH($A{r},{R["종목"]},0))', align=CENTER)
        for j in range(3):
            cl = get_column_letter(3 + j)
            s.put(f"{cl}{r}", f'=SUMIFS({R["금액"]},{R["종목"]},$A{r},{R["주차"]},{CW},{R["주체"]},{cl}${r1 + 2})', fmt=NUM)
        s.put(f"F{r}", f"=SUM(C{r}:E{r})", "b", NUM)
        s.put(f"G{r}", f"=COUNTIFS({R['종목']},$A{r},{R['주차']},{CW})", align=CENTER)
        s.put(f"H{r}", f"=_xlfn.MAXIFS({R['연속']},{R['종목']},$A{r},{R['주차']},{CW})", align=CENTER)
        s.put(f"I{r}", f"=AVERAGEIFS({R['등락']},{R['종목']},$A{r},{R['주차']},{CW})", fmt=PCT)
        s.put(f"J{r}", f'=IF(COUNTIFS({R["종목"]},$A{r},{R["주차"]},{PW})>0,"O","신규")', align=CENTER)
        s.put(f"K{r}", f'=TRIM(IF(G{r}>=3,"3주체동시 ",IF(G{r}=2,"쌍끌이 ",""))&IF(J{r}="신규","신규진입 ","")'
                       f'&IF(AND(H{r}>={PULLBACK_STREAK},I{r}<0),"눌림매집 ","")&IF(H{r}>={LONG_STREAK},"장기연속",""))', "b")
    s_last = s_first + len(order) - 1
    cf = ws.conditional_formatting
    cf.add(f"I{s_first}:I{s_last}", CellIsRule(operator="lessThan", formula=["0"], font=Font(color="C00000")))
    cf.add(f"I{s_first}:I{s_last}", CellIsRule(operator="greaterThanOrEqual", formula=["10"], font=Font(color="0B6E2E", bold=True)))
    cf.add(f"A{s_first}:K{s_last}", FormulaRule(formula=[f"$G{s_first}>=3"], fill=PatternFill("solid", fgColor="FFF2CC")))
    cf.add(f"A{s_first}:K{s_last}", FormulaRule(formula=[f"AND($H{s_first}>={PULLBACK_STREAK},$I{s_first}<0)"],
                                                fill=PatternFill("solid", fgColor="E2EFDA")))
    s.put(f"A{s_last + 1}", "노란 행 = 3주체 동시 매수 · 연두 행 = 눌림매집 후보 · 등락률 빨강 = 하락, 초록 굵게 = +10% 이상",
          "note", border=False)

    # 표 3: 최근 WINDOW주 누적 순위
    wdf = df[df["주차"].isin(win)]
    top = wdf.groupby("종목명")["순매수금액"].sum().sort_values(ascending=False).head(TOP_N).index.tolist()
    r2 = s_last + 4
    s.section(r2, f"3. 최근 {len(win)}주 누적 순매수 상위 {len(top)}종목 ({win[0]} ~ {cur})",
              f"등장 주수가 많을수록 꾸준한 매수. 참여주체 3이면 외국인·기관·연기금 모두 산 종목. 기간은 표 1의 마지막 {len(win)}주")
    s.header(r2 + 2, 1, ["순위", "종목명", "외국인", "기관합계", "연기금", "합계", f"등장 주수(/{len(win)})",
                         "참여주체수", "최대 연속순매수일", "이번주 등락률(%)"])
    t_first = r2 + 3
    for i, name in enumerate(top):
        r = t_first + i
        s.put(f"A{r}", i + 1, align=CENTER)
        s.put(f"B{r}", name, "b")
        for j in range(3):
            cl = get_column_letter(3 + j)
            s.put(f"{cl}{r}", f'=SUMPRODUCT(SUMIFS({R["금액"]},{R["종목"]},$B{r},{R["주체"]},{cl}${r2 + 2},{R["주차"]},{WIN}))', fmt=NUM)
        s.put(f"F{r}", f"=SUM(C{r}:E{r})", "b", NUM)
        s.put(f"G{r}", f"=SUMPRODUCT(--(COUNTIFS({R['종목']},$B{r},{R['주차']},{WIN})>0))", align=CENTER)
        s.put(f"H{r}", f"=(C{r}>0)+(D{r}>0)+(E{r}>0)", align=CENTER)
        s.put(f"I{r}", f"=SUMPRODUCT(MAX(({R['종목']}=$B{r})*ISNUMBER(MATCH({R['주차']},{WIN},0))*{R['연속']}))", align=CENTER)
        s.put(f"J{r}", f'=IFERROR(AVERAGEIFS({R["등락"]},{R["종목"]},$B{r},{R["주차"]},{CW}),"-")', fmt=PCT,
              align=Alignment(horizontal="right"))
    t_last = t_first + len(top) - 1
    cf.add(f"F{t_first}:F{t_last}", DataBarRule(start_type="min", end_type="max", color="5B9BD5"))
    s.put(f"A{t_last + 1}", "이번주 등락률 '-' = 이번 주 상위 10에 없음", "note", border=False)

    # 누적 차트: 1위가 2위의 3배를 넘으면 축 왜곡을 막으려 1위를 빼고 그린다
    tot = wdf.groupby("종목명")["순매수금액"].sum()
    skip = 1 if len(top) > 2 and tot[top[0]] > 3 * tot[top[1]] else 0
    n_chart = min(15, len(top) - skip)
    ch2 = BarChart()
    ch2.type, ch2.grouping, ch2.overlap = "bar", "stacked", 100
    ch2.title = f"최근 {len(win)}주 누적 순매수 (억원" + (f", 1위 {top[0]} 제외)" if skip else ")")
    for j in range(3):
        ser = Series(Reference(ws, min_col=3 + j, min_row=t_first + skip, max_row=t_first + skip + n_chart - 1))
        ser.tx = SeriesLabel(strRef=StrRef(f"'{OUT_SHEET}'!${get_column_letter(3 + j)}${r2 + 2}"))
        ch2.series.append(ser)
    ch2.set_categories(Reference(ws, min_col=2, min_row=t_first + skip, max_row=t_first + skip + n_chart - 1))
    ch2.x_axis.scaling.orientation = "maxMin"
    ch2.x_axis.tickLblSkip = 1
    ch2.y_axis.numFmt = "#,##0"
    ch2.legend.position = "b"
    ch2.gapWidth = 40
    ch2.height, ch2.width = 13, 17
    ch2.x_axis.delete = ch2.y_axis.delete = False
    ws.add_chart(ch2, f"L{r2}")

    # 표 4: 시장별
    r3 = t_last + 4
    s.section(r3, "4. 시장별 · 투자주체별 전체 기간 누적 (억원)", "주체별로 대형주(KOSPI)와 중소형주(KOSDAQ) 중 어디를 샀는지")
    s.header(r3 + 2, 1, ["시장"] + INV + ["합계"])
    for i, mkt in enumerate(["KOSPI", "KOSDAQ"]):
        r = r3 + 3 + i
        s.put(f"A{r}", mkt, "b", align=CENTER)
        for j in range(3):
            cl = get_column_letter(2 + j)
            s.put(f"{cl}{r}", f'=SUMIFS({R["금액"]},{R["시장"]},$A{r},{R["주체"]},{cl}${r3 + 2})', fmt=NUM)
        s.put(f"E{r}", f"=SUM(B{r}:D{r})", "b", NUM)
    rk, rsh = r3 + 4, r3 + 5
    s.put(f"A{rsh}", "KOSDAQ 비중", "b", fill=FILL_TOT, align=CENTER)
    for cl in "BCDE":
        s.put(f"{cl}{rsh}", f"=IFERROR({cl}{rk}/({cl}{rk - 1}+{cl}{rk}),0)", "b", "0.0%", FILL_TOT)

    for k, v in {"A": 15, "B": 15, "C": 11, "D": 11, "E": 11, "F": 11, "G": 11, "H": 11,
                 "I": 12, "J": 11, "K": 20, "N": 11}.items():
        ws.column_dimensions[k].width = v
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    if themes:
        write_theme_sheet(wb, themes, cur, len(win))
    wb.active = 0
    wb.calculation.fullCalcOnLoad = True  # 엑셀이 열 때 항상 다시 계산
    wb.save(path)

    # 노션·메일용 요약 데이터 (같은 규칙을 파이썬으로 계산)
    g = cdf.groupby("종목명").agg(
        시장=("시장", "first"), 참여=("투자주체", "nunique"), 연속=("연속순매수일", "max"), 등락=("주간등락률", "mean"))
    amt = cdf.pivot_table(index="종목명", columns="투자주체", values="순매수금액", aggfunc="sum", fill_value=0)
    prev_names = set(df[df["주차"] == prev]["종목명"]) if prev != cur else set()
    signals = []
    for name in order:
        row = g.loc[name]
        tags = []
        if row.참여 >= 3:
            tags.append("3주체동시")
        elif row.참여 == 2:
            tags.append("쌍끌이")
        if name not in prev_names and prev != cur:
            tags.append("신규진입")
        if row.연속 >= PULLBACK_STREAK and row.등락 < 0:
            tags.append("눌림매집")
        if row.연속 >= LONG_STREAK:
            tags.append("장기연속")
        signals.append({
            "종목명": name, "시장": row.시장,
            **{inv: float(amt.loc[name, inv]) if inv in amt else 0.0 for inv in INV},
            "참여": int(row.참여), "연속": int(row.연속), "등락": round(float(row.등락), 2), "신호": " ".join(tags),
        })
    cum = wdf.pivot_table(index="종목명", columns="투자주체", values="순매수금액", aggfunc="sum", fill_value=0)
    cum["합계"] = cum.sum(axis=1)
    cum = cum.sort_values("합계", ascending=False).head(10)
    return {
        "week": cur, "title": week_title(cur), "filename": report_filename(cur),
        "weeks": weeks, "window": win, "insights": insights, "signals": signals,
        "cumulative": [{"종목명": n, **{c: float(cum.loc[n, c]) for c in cum.columns}} for n in cum.index],
        "themes": themes,
    }
